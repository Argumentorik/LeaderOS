"""One-off: Word-Übersicht für Wlad generieren + per Resend an wj@argumentorik.com senden."""
import base64
import os
import sys

from dotenv import load_dotenv
load_dotenv("/app/backend/.env")

from docx import Document
from docx.shared import Pt, RGBColor

LIME = RGBColor(0x6B, 0x8A, 0x00)
DARK = RGBColor(0x0A, 0x0A, 0x0A)

ANMELDUNG = [
    ("1. Login ohne Passwort (Magic Link)",
     "Gehe auf leaderos.de/login, klicke oben auf den Tab \u201eMagic Link\u201c, gib wj@argumentorik.com ein und klicke auf \u201eLink senden\u201c. Du bekommst innerhalb von Sekunden eine E-Mail mit einem sicheren 1-Klick-Login-Link \u2014 kein Passwort n\u00f6tig. Alternativ funktioniert auch \u201eMit Google fortfahren\u201c."),
    ("2. Dein Account hat automatisch Admin-Rechte",
     "Sobald du dich mit wj@argumentorik.com einloggst, wird dein Konto automatisch mit vollen Admin-Rechten ausgestattet. Du landest zuerst auf deinem pers\u00f6nlichen Dashboard mit WladBot 3.0, dem 30-Tage-Sprint und allen Features."),
    ("3. Admin-Bereich \u00f6ffnen",
     "Der Admin-Bereich liegt bewusst auf einer nicht \u00f6ffentlichen URL: leaderos.de/wlad-control-x7k9q2 \u2014 dort siehst du Umsatz (MRR/ARR), alle User, Transaktionen, Sync-Status und ganz unten die Lizenz-Key-Verwaltung."),
]

LIZENZ = [
    "Im Admin-Panel ganz nach unten scrollen zur Sektion \u201eLizenz-Keys \u00b7 Sales-Aktivierung\u201c.",
    "Paket ausw\u00e4hlen: \u201eLeadership OS (Standard \u00b7 997\u20ac)\u201c oder \u201eLeadership OS PLUS (4.797\u20ac)\u201c.",
    "Anzahl eingeben (1\u2013100) und optional eine Notiz (z. B. Kundenname oder \u201eMasterclass Batch 1\u201c) \u2014 dann auf \u201eGenerieren\u201c klicken.",
    "WICHTIG: Die Keys werden aus Sicherheitsgr\u00fcnden nur EIN einziges Mal im Klartext angezeigt. Sofort auf \u201eAlle kopieren\u201c klicken und sicher ablegen (z. B. in eurer Sales-Tabelle).",
    "Key an den Kunden geben (telefonisch, per E-Mail oder im Sales-Call). Format: XXXX-XXXX-XXXX-XXXX \u2014 16 Zeichen, ohne verwechselbare Zeichen wie 0/O oder 1/I.",
    "Der Kunde l\u00f6st ein: leaderos.de/login \u2192 Tab \u201eLizenz-Key\u201c \u2192 Key + Name + E-Mail + Passwort eingeben \u2192 Paket wird SOFORT freigeschaltet (neues Konto wird automatisch erstellt, bestehende Konten werden upgegradet).",
    "In der Key-Liste siehst du jederzeit den Status: aktiv / eingel\u00f6st (inkl. Kunden-E-Mail) / deaktiviert. Falsch ausgegebene Keys kannst du mit dem Sperr-Symbol deaktivieren.",
]

HIGHLIGHTS = [
    ("\U0001F916 WladBot 3.0 \u2014 dein digitaler Zwilling",
     "Ein KI-Coach, der deine komplette Methodik kennt: alle Kurse, Frameworks und Inhalte sind als durchsuchbares Wissen hinterlegt (RAG mit Voyage-Embeddings). Antwortet 24/7 im Stil von Wlad \u2014 mit Coaching-Modus, Chat-Verlauf und Ordnern wie bei ChatGPT."),
    ("\U0001F3AF 30-Tage KI-Leadership Sprint",
     "300 Fragen, t\u00e4gliche Challenges und echtes Momentum: Jeden Tag eine Aufgabe, die den Teilnehmer als Leader sch\u00e4rft \u2014 gamifiziert mit XP, Levels und Fortschritts-Tracking."),
    ("\U0001F3A5 16 Video-Missionen mit KI-Analyse",
     "Kamera an, 60 Sekunden Rede \u2014 die KI bewertet nach deinem Framework in Sekunden: Klarheit, Souver\u00e4nit\u00e4t, Empathie, Struktur. Mit Transkript, pers\u00f6nlichem Wlad-Feedback und Download-Report."),
    ("\U0001F94A KI-Simulations-Arena",
     "Realistische Roleplays f\u00fcr schwierige F\u00fchrungssituationen (K\u00fcndigungsgespr\u00e4ch, Gehaltsverhandlung, Konflikt im Team). Mit Briefing-Wizard: Gegen\u00fcber-Typ und Trainingsziel w\u00e4hlbar."),
    ("\U0001F4CA Leadership-Diagnose + leadercheck.de Sync",
     "Wer auf leadercheck.de den Test macht, bekommt seine Scores automatisch ins Leader-OS synchronisiert \u2014 WladBot kennt die Ergebnisse und coacht gezielt auf die Schw\u00e4chen."),
    ("\U0001F4B3 Stripe Live-Payments mit Ratenzahlung",
     "4 Pakete live und getestet: Leadership OS (997\u20ac), 2\u00d7-Rate, 12\u00d7-Rate (99\u20ac/Monat) und PLUS (4.797\u20ac). Automatische Freischaltung, Rechnungs- und Willkommens-Mails inklusive."),
    ("\U0001F39F\uFE0F Lizenz-Key-System f\u00fcr Sales",
     "16-stellige Aktivierungscodes, die du selbst im Admin-Panel generierst \u2014 perfekt f\u00fcr telefonisch verkaufte Pakete und die 400 Masterclass-Kunden. Einl\u00f6sung in unter 60 Sekunden auf der Login-Seite."),
    ("\U0001F3C6 Gamification, Community & Live-Events",
     "XP-System, Level, Empfehlungsprogramm, Community-Feed und Live-KI-Calls mit Wlad (Event-Season ab Mitte August) \u2014 inkl. Kalender-Integration und Erinnerungs-Mails."),
    ("\U0001F6E1\uFE0F Enterprise-Security \u2014 4 bestandene Audits",
     "Magic-Link-Login, Brute-Force-Schutz, Ger\u00e4te-Erkennung mit Sicherheits-Mail bei neuem Login, Passwort-Mindestst\u00e4rke, DSGVO-konform. Der letzte Security-Audit kam ohne einen einzigen Befund zur\u00fcck."),
    ("\U0001F4C8 Admin-Cockpit f\u00fcr den vollen \u00dcberblick",
     "Umsatz in Echtzeit (MRR/ARR), User- und Tier-Verwaltung, Transaktionen, Sync-Health, RAG-Wissens-Studio (zeigt, welche deiner Inhalte WladBot pro Frage nutzt) und die Lizenz-Key-Verwaltung."),
]


def build_docx(path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    h = doc.add_heading("Leader-OS \u00b7 Dein Zugang, Lizenz-Keys & alles, was wir gebaut haben", level=0)
    for run in h.runs:
        run.font.color.rgb = DARK
    p = doc.add_paragraph("F\u00fcr Wlad \u00b7 Stand: Juli 2026 \u00b7 leaderos.de")
    p.runs[0].font.color.rgb = LIME
    p.runs[0].font.bold = True

    doc.add_heading("1. So meldest du dich an \u2014 in 3 Schritten", level=1)
    for title, body in ANMELDUNG:
        doc.add_heading(title, level=2)
        doc.add_paragraph(body)

    doc.add_heading("2. Lizenz-Keys selbst erstellen & vergeben", level=1)
    doc.add_paragraph(
        "Mit Lizenz-Keys schaltest du Kunden frei, die NICHT \u00fcber Stripe kaufen \u2014 z. B. Sales-Deals am Telefon "
        "oder die Masterclass-Teilnehmer. Ein Key = ein Kunde = ein Paket."
    )
    for i, step in enumerate(LIZENZ, 1):
        doc.add_paragraph(f"{i}. {step}")

    doc.add_heading("3. Was wir gebaut haben \u2014 10 Highlights", level=1)
    for title, body in HIGHLIGHTS:
        doc.add_heading(title, level=2)
        doc.add_paragraph(body)

    doc.add_paragraph()
    footer = doc.add_paragraph("Fragen? Mert erreicht das Team jederzeit \u2014 und WladBot beantwortet den Rest. \U0001F680")
    footer.runs[0].font.italic = True
    doc.save(path)


EMAIL_HTML = """
<div style="max-width:640px;margin:0 auto;font-family:-apple-system,Segoe UI,Helvetica,Arial,sans-serif;background:#0A0A0A;color:#f5f5f5;border-radius:16px;overflow:hidden">
  <div style="padding:28px 32px;border-bottom:1px solid rgba(255,255,255,0.08)">
    <p style="margin:0;font-size:13px;letter-spacing:2px;color:#BFFF00;font-weight:700">LEADER\u00b7OS &middot; POWERED BY WLADBOT 3.0</p>
    <h1 style="margin:10px 0 0;font-size:24px;color:#ffffff">Hallo Wlad \u2014 dein Zugang steht! \U0001F680</h1>
    <p style="margin:8px 0 0;font-size:14px;color:rgba(255,255,255,0.6)">Alles Wichtige in dieser Mail \u2014 die ausf\u00fchrliche Version findest du zus\u00e4tzlich als Word-Datei im Anhang.</p>
  </div>

  <div style="padding:28px 32px">
    <h2 style="font-size:17px;color:#BFFF00;margin:0 0 14px">\u2460 So meldest du dich an \u2014 3 Schritte</h2>
    <ol style="margin:0;padding-left:20px;font-size:14px;line-height:1.7;color:rgba(255,255,255,0.85)">
      <li style="margin-bottom:10px"><strong style="color:#fff">Magic Link nutzen:</strong> Auf <a href="https://leaderos.de/login" style="color:#BFFF00">leaderos.de/login</a> den Tab <strong>\u201eMagic Link\u201c</strong> w\u00e4hlen, <strong>wj@argumentorik.com</strong> eingeben \u2192 du bekommst sofort einen 1-Klick-Login-Link per Mail. Kein Passwort n\u00f6tig. (Alternativ: \u201eMit Google fortfahren\u201c.)</li>
      <li style="margin-bottom:10px"><strong style="color:#fff">Admin-Rechte automatisch:</strong> Dein Konto (wj@argumentorik.com) wird beim ersten Login automatisch mit vollen Admin-Rechten ausgestattet.</li>
      <li><strong style="color:#fff">Admin-Bereich \u00f6ffnen:</strong> <a href="https://leaderos.de/wlad-control-x7k9q2" style="color:#BFFF00">leaderos.de/wlad-control-x7k9q2</a> \u2014 dein Cockpit mit Umsatz, Usern &amp; Lizenz-Keys.</li>
    </ol>

    <h2 style="font-size:17px;color:#BFFF00;margin:28px 0 14px">\u2461 Lizenz-Keys selbst erstellen &amp; vergeben</h2>
    <div style="background:rgba(191,255,0,0.06);border:1px solid rgba(191,255,0,0.2);border-radius:12px;padding:16px 18px;font-size:14px;line-height:1.7;color:rgba(255,255,255,0.85)">
      <p style="margin:0 0 8px"><strong style="color:#fff">Generieren:</strong> Im Admin-Panel ganz unten \u2192 Sektion <strong>\u201eLizenz-Keys \u00b7 Sales-Aktivierung\u201c</strong> \u2192 Paket w\u00e4hlen (997\u20ac oder PLUS 4.797\u20ac), Anzahl + Notiz eingeben \u2192 <strong>\u201eGenerieren\u201c</strong>.</p>
      <p style="margin:0 0 8px"><strong style="color:#fff">\u26A0\uFE0F Wichtig:</strong> Keys sind nur <u>einmal</u> sichtbar \u2192 sofort \u201eAlle kopieren\u201c und sicher ablegen.</p>
      <p style="margin:0"><strong style="color:#fff">Kunde l\u00f6st ein:</strong> <a href="https://leaderos.de/login" style="color:#BFFF00">leaderos.de/login</a> \u2192 Tab \u201eLizenz-Key\u201c \u2192 Key + E-Mail + Passwort \u2192 Paket sofort aktiv. Status (aktiv/eingel\u00f6st) siehst du jederzeit in der Liste \u2014 Deaktivieren per Klick m\u00f6glich.</p>
    </div>

    <h2 style="font-size:17px;color:#BFFF00;margin:28px 0 14px">\u2462 Was wir gebaut haben \u2014 10 Highlights</h2>
    <ul style="margin:0;padding-left:2px;list-style:none;font-size:13.5px;line-height:1.65;color:rgba(255,255,255,0.85)">
      __HIGHLIGHTS__
    </ul>

    <p style="margin:28px 0 0;font-size:13px;color:rgba(255,255,255,0.5)">\U0001F4CE Im Anhang: <strong style="color:rgba(255,255,255,0.75)">Leader-OS_Uebersicht_Wlad.docx</strong> \u2014 alles nochmal ausf\u00fchrlich zum Nachlesen &amp; Weiterleiten.</p>
  </div>

  <div style="padding:18px 32px;border-top:1px solid rgba(255,255,255,0.08);font-size:11px;color:rgba(255,255,255,0.4)">
    Leader\u00b7OS \u00b7 Argumentorik GmbH \u00b7 <a href="https://leaderos.de" style="color:rgba(191,255,0,0.7)">leaderos.de</a>
  </div>
</div>
"""


def main():
    docx_path = "/tmp/Leader-OS_Uebersicht_Wlad.docx"
    build_docx(docx_path)
    print("DOCX created:", os.path.getsize(docx_path), "bytes")

    hl_html = "".join(
        f'<li style="margin-bottom:9px"><strong style="color:#fff">{t}</strong><br/><span style="color:rgba(255,255,255,0.65)">{b}</span></li>'
        for t, b in HIGHLIGHTS
    )
    html = EMAIL_HTML.replace("__HIGHLIGHTS__", hl_html)

    import resend
    resend.api_key = os.environ["RESEND_API_KEY"]
    sender = os.environ.get("SENDER_EMAIL", "onboarding@resend.dev")
    with open(docx_path, "rb") as f:
        content_b64 = base64.b64encode(f.read()).decode()

    result = resend.Emails.send({
        "from": f"Leader-OS <{sender}>",
        "to": [sys.argv[1] if len(sys.argv) > 1 else "wj@argumentorik.com"],
        "subject": "Dein Leader-OS Zugang + Lizenz-Keys \u2014 alles in 3 Schritten (inkl. Word-Datei) \U0001F680",
        "html": html,
        "attachments": [{"filename": "Leader-OS_Uebersicht_Wlad.docx", "content": content_b64}],
    })
    print("Resend result:", result)


if __name__ == "__main__":
    main()
